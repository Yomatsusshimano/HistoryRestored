"""Read-only NARA Special List25 locator audit; requires openpyxl.

Run from repository root: python analysis/willapa_nara_finding_aid.py
Preserves spellings, missing fields and coded years without interpreting them.
"""
from pathlib import Path
import hashlib
import json
import openpyxl

SOURCE=Path('sources/originals/cascadia/Special-List-25-United-States.xlsx')


def audit():
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()=='135ba9b2bca4c068357555ef067df8a97a319810854fcc851c57e990ad23a5de', 'Unexpected finding-aid version'
    wb=openpyxl.load_workbook(SOURCE,read_only=True,data_only=False)
    ws=wb['Washington']
    header=next(ws.iter_rows(min_row=1,max_row=1,values_only=True))
    assert header==('County/Area','Symbol','Year','RG','# of Indexes','Index Type','Scale of Negatives','Filed Under','IM/NUS Number','State')
    selected=[]
    for i,values in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
        if values[0] in ('Pacific','Willpa Bay','Willapa Bay'):
            selected.append({'sheet':'Washington','row':i,'range':f'A{i}:J{i}',
                             'source_values':dict(zip(header,values))})
    assert [r['row'] for r in selected]==[340,341,342,530]
    assert selected[-1]['source_values']['County/Area']=='Willpa Bay'
    assert selected[-1]['source_values']['RG']=='23'
    assert selected[-1]['source_values']['Filed Under']=='Project Completion Reports'
    formulas=[{'row':i,'column':c.column_letter,'formula':c.value}
              for i in (1,340,341,342,530) for c in ws[i] if c.data_type=='f']
    assert not formulas, 'Selected locators unexpectedly contain formulas'
    wb.close()
    return {'source_id':'S268','accessed_date':'2026-10-09',
            'url':'https://unwritten-record.blogs.archives.gov/wp-content/uploads/sites/6/2020/01/Special-List-25-United-States.xlsx',
            'local_copy':SOURCE.as_posix(),'bytes':SOURCE.stat().st_size,
            'sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),
            'inspected_scope':'Washington sheet headers and all546rows; selected four locator rows; other sheets not audited for holdings',
            'rows':selected,'selected_formula_cells':formulas,
            'limits':['Willpa Bay is preserved as written; interpreting it as Willapa is a retrieval hypothesis.',
                      'No exact target frame, annotated enlargement, field sheet or accession is authenticated.',
                      '1949P,1951P,****P,LI,PI-L and asterisks retained literally; code meanings not assumed.',
                      'County-level entries do not establish coverage of Grassy Island.']}


if __name__=='__main__':
    result=audit()
    Path('data/willapa-nara-finding-aid.json').write_text(json.dumps(result,indent=2)+'\n',encoding='utf-8')
    print('NARA finding-aid audit: four locator rows preserved; no original target image recovered.')
