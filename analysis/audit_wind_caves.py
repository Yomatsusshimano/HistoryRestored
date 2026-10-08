"""Crosswalk published 04PW30 ages; no raw isotope reduction or bed dating."""
import argparse, hashlib, json, math
from pathlib import Path
import openpyxl
p=argparse.ArgumentParser();p.add_argument('directory',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
def sheet(n,name):
    return openpyxl.load_workbook(a.directory/f'G48080_SuppTab{n}.xlsx',read_only=True,data_only=True)[name]
s=sheet(2,'summary data'); d=sheet(2,'data'); t=sheet(4,'Sheet1')
rows=[]
for r,v in enumerate(d.iter_rows(values_only=True),1):
    if len(v)>20 and v[20]=='04PW30':
        rows.append({'data_row':r,'run_id':v[19],'age_Ma':v[9],'one_sigma_Ma':v[10]})
selected=[]
for r in range(66,71):
    age,error=s.cell(r,3).value,s.cell(r,4).value
    matches=[x for x in rows if x['age_Ma']==age and x['one_sigma_Ma']==error]
    assert len(matches)==1
    selected.append(dict(matches[0],summary_row=r))
w=sum(1/x['one_sigma_Ma']**2 for x in selected)
mean=sum(x['age_Ma']/x['one_sigma_Ma']**2 for x in selected)/w
mswd=sum(((x['age_Ma']-mean)/x['one_sigma_Ma'])**2 for x in selected)/(len(selected)-1)
error=2/math.sqrt(w)
assert len(rows)==57 and len({x['run_id'] for x in rows})==57
assert abs(mean-s['C71'].value)<1e-12 and abs(error-s['D71'].value)<1e-12
assert round(mswd,2)==s['C72'].value
loc={ 'table':4,'sheet':'Sheet1','row':6,'sample':t['A6'].value,'unit':t['B6'].value,'reported_maximum_depositional_age':t['C6'].value,'latitude':t['D6'].value,'longitude':t['E6'].value,'datum':'NAD83','elevation':None,'bed_height':None,'status':'source-reported, not independently surveyed'}
result={'source_id':'S210','sample':'04PW30','scope':'Stored analytical ages and published selected-group arithmetic; shared calibration covariance and raw isotope reduction not reproduced.','hashes':{f'G48080_SuppTab{n}.xlsx':hashlib.sha256((a.directory/f'G48080_SuppTab{n}.xlsx').read_bytes()).hexdigest() for n in (2,4)},'location':loc,'data_header':d['B1211'].value,'all_57_rows':rows,'selected_five':selected,'reported':{'mean_Ma':s['C71'].value,'two_sigma_Ma':s['D71'].value,'MSWD':s['C72'].value,'selection':s['C73'].value},'calculated':{'mean_Ma':mean,'two_sigma_internal_Ma':error,'MSWD':mswd},'next_oldest_row':sorted(rows,key=lambda x:x['age_Ma'])[5],'interpretation':'Conditional maximum depositional age: host deposition no older than valid youngest detrital crystals; not a minimum bed age or proof of a historical event.'}
a.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(result['calculated']));print(result['location'])
