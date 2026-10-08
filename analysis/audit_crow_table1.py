"""Read-only audit of Crow2021 Table1; requires openpyxl. No Excel execution."""
import argparse, hashlib, json, math
from pathlib import Path
import openpyxl
p=argparse.ArgumentParser()
p.add_argument('workbook',type=Path)
p.add_argument('output',type=Path)
a=p.parse_args()
w=openpyxl.load_workbook(a.workbook,data_only=False)
s=w['Sheet1']
rows=[]
for r in range(20,25):
 rows.append({'row':r,'cells':{c.coordinate:{'value':c.value,'underline':c.font.underline,'type':c.data_type} for c in s[r] if c.value is not None}})
x=[(s.cell(r,8).value,s.cell(r,9).value) for r in (20,21)]
result={'source_id':'S210','sha256':hashlib.sha256(a.workbook.read_bytes()).hexdigest(),'inspection':'Stored cells and font metadata; workbook not rendered or executed.','lawlor_rows':rows,'conditional_two_row_calculation':{'assumption':'Independent displayed ages/errors, inverse variance weighting; no shared calibration covariance or additional inputs. Errors follow table 2sigma convention.','mean_Ma':sum(v/e**2 for v,e in x)/sum(1/e**2 for v,e in x),'internal_2sigma_Ma':1/math.sqrt(sum(1/e**2 for v,e in x))},'auxiliary_broken_references':[{'cell':c.coordinate,'formula':c.value} for row in w['Sheet2'] for c in row if c.data_type=='f' and '#REF!' in c.value]}
a.output.write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
print('Preserved five Lawlor rows and',len(result['auxiliary_broken_references']),'broken references')
