"""Read-only audit of source XLSX stored/cached values versus source CSV."""
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import zipfile
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
ledger=json.loads((ROOT/'data/cascadia-field-acquisition.json').read_text())
def archive(name):
    entry=next(f for f in ledger['files'] if f['name']==name)
    raw=(ROOT/entry['file']).read_bytes()
    assert hashlib.sha256(raw).hexdigest()==entry['sha256']
    return zipfile.ZipFile(io.BytesIO(raw)),entry['sha256']
plain,plain_hash=archive('plain.zip')
formatted_hash=next(f['sha256'] for f in ledger['files'] if f['name']=='supplement.zip')
selected=['03_Soil_stratigraphy','11_Redcedar_dead_sampled',
          '13_Redcedar_live','18_Copalis_vented_sand']
out={'status':'SOURCE_FORMAT_COMPARISON_NOT_RECALCULATION_OR_SCIENTIFIC_VALIDATION',
     'source_id':'S244','plain_zip_sha256':plain_hash,
     'supplement_zip_sha256':formatted_hash,'tables':[]}
for stem in selected:
    p=ROOT/'sources/originals/cascadia/field-release-2022'/(stem+'.xlsx')
    xbytes=p.read_bytes()
    entry=next(f for f in ledger['formatted_members'] if f['file']==p.relative_to(ROOT).as_posix())
    assert hashlib.sha256(xbytes).hexdigest()==entry['sha256']
    wb=openpyxl.load_workbook(io.BytesIO(xbytes),data_only=False)
    cached=openpyxl.load_workbook(io.BytesIO(xbytes),data_only=True)
    sheet=wb.active
    cache=cached[sheet.title]
    csv_rows=list(csv.reader(io.StringIO(plain.read(stem+'.csv').decode('cp1252'))))
    headers=csv_rows[0]
    named=[i for i,h in enumerate(headers) if h]
    assert [sheet.cell(1,i+1).value for i in named]==[headers[i] for i in named]
    assert sheet.max_row==len(csv_rows)
    differences=[]
    formulas=[]
    numeric=0
    field_rows=[]
    for rownum,row in enumerate(csv_rows[1:],2):
        if not any(row[i].strip() for i in named):
            continue
        identity={headers[i]:row[i] for i in named
                  if headers[i] in ['FieldID','TreeID','LocSeq','Category']}
        values={headers[i]:cache.cell(rownum,i+1).value for i in named}
        field_rows.append({'xlsx_row':rownum,'csv_row':rownum,
                           'identity':identity,'stored_values':values})
        for i in named:
            cell=sheet.cell(rownum,i+1)
            value=cache.cell(rownum,i+1).value
            if cell.data_type=='f':
                formulas.append({'cell':cell.coordinate,'formula':cell.value,
                                 'cached_value':value,'csv_value':row[i],
                                 'number_format':cell.number_format})
            if isinstance(value,(int,float)) and not isinstance(value,bool):
                numeric+=1
                try:
                    plain_num=float(row[i])
                except ValueError:
                    differences.append({'cell':cell.coordinate,'field':headers[i],
                                        'identity':identity,'stored_value':value,
                                        'csv_value':row[i],'kind':'NUMERIC_VS_NONNUMERIC',
                                        'number_format':cell.number_format})
                    continue
                if not math.isclose(float(value),plain_num,rel_tol=0,abs_tol=1e-12):
                    differences.append({'cell':cell.coordinate,'field':headers[i],
                                        'identity':identity,'stored_value':value,
                                        'csv_value':row[i],'kind':'NUMERIC_VALUE_DIFFERENCE',
                                        'number_format':cell.number_format})
    out['tables'].append({'name':stem,'sheet':sheet.title,
                          'xlsx_file':p.relative_to(ROOT).as_posix(),
                          'xlsx_sha256':hashlib.sha256(xbytes).hexdigest(),
                          'csv_sha256':hashlib.sha256(plain.read(stem+'.csv')).hexdigest(),
                          'nonblank_rows':len(field_rows),'numeric_cells_compared':numeric,
                          'numeric_difference_count':len(differences),
                          'numeric_differences':differences,'formula_cells':formulas,
                          'source_rows':field_rows})
soil=out['tables'][0]
assert next(r for r in soil['source_rows'] if r['xlsx_row']==4)['stored_values']['Depth']==0.55
for field,depths in [('T6',[0.4,0.9,0.4,0.9]),('T1',[1.3,1.8,0.3]),('A31',[0.9,0.5,0.8,0.9])]:
    assert [r['stored_values']['Depth'] for r in soil['source_rows']
            if r['stored_values']['FieldID']==field]==depths
out['limits']=['Original workbooks remain unchanged.',
               'Stored values and formula caches were read; no spreadsheet-engine recalculation.',
               'Extra stored decimal places are not independently established field accuracy.',
               'Date/text formatting and every annotation were not exhaustively compared.',
               'Source versions are dependent; agreement is not independent observation.']
(ROOT/'data/cascadia-format-comparison.json').write_text(json.dumps(out,indent=2,default=str)+'\n')
for t in out['tables']:
    print(t['name'],t['nonblank_rows'],'rows;',t['numeric_difference_count'],
          'numeric differences;',len(t['formula_cells']),'formula cells')
print('Depth inversions persist in XLSX stored values')
