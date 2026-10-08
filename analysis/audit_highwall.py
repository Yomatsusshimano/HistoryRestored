"""Extract stored Highwall Wash results; reproduce selected-row arithmetic only."""
import argparse,json,math,hashlib,re
from pathlib import Path
from collections import Counter
import openpyxl
p=argparse.ArgumentParser();p.add_argument('directory',type=Path);p.add_argument('output',type=Path);a=p.parse_args()
def sheet(n,title):return openpyxl.load_workbook(a.directory/f'G48080_SuppTab{n}.xlsx',read_only=True,data_only=True)[title]
s=sheet(2,'summary data');rows=[{'row':r,'age_Ma':s.cell(r,3).value,'one_sigma_Ma':s.cell(r,4).value} for r in range(24,43)]
weight=sum(1/x['one_sigma_Ma']**2 for x in rows);mean=sum(x['age_Ma']/x['one_sigma_Ma']**2 for x in rows)/weight
mswd=sum(((x['age_Ma']-mean)/x['one_sigma_Ma'])**2 for x in rows)/(len(rows)-1)
t=sheet(6,'Sheet1');mag=[]
for r in range(3,39):
 vals=[t.cell(r,c).value for c in range(1,8)]
 mag.append({'row':r,'sample':vals[0].strip(),'demagnetization':vals[1],'polarity':vals[2],'declination':vals[3],'inclination':vals[4],'MAD':vals[5],'rank':vals[6].strip()})
result={'source_id':'S210','scope':'Stored selected age rows and all Table6 specimen rows. No raw isotope or magnetic-vector refit.','files':{str(n):hashlib.sha256((a.directory/f'G48080_SuppTab{n}.xlsx').read_bytes()).hexdigest() for n in (2,6)},'selected_ages':rows,'reported':{'mean_Ma':s['C43'].value,'two_sigma_Ma':s['D43'].value,'MSWD':s['C44'].value,'selection':s['C45'].value},'calculated':{'mean_Ma':mean,'MSWD':mswd,'two_sigma_internal_Ma':2/math.sqrt(weight),'two_sigma_scatter_scaled_Ma':2*math.sqrt(max(1,mswd)/weight)},'magnetic_specimens':mag,'polarity_counts':dict(Counter(x['polarity'] for x in mag)),'table6_note':t['A39'].value,'coordinate_status':'Source-reported only; not independently geolocated or corrected.'}
a.output.write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
assert abs(mean-result['reported']['mean_Ma'])<1e-12
assert abs(result['calculated']['two_sigma_scatter_scaled_Ma']-result['reported']['two_sigma_Ma'])<1e-12
assert round(mswd,2)==result['reported']['MSWD']
print(result['calculated']);print(result['polarity_counts'])
