"""Read-only extraction and descriptive comparisons, without correlation thresholds."""
import hashlib
import json
import sys
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
original = Path(sys.argv[1])
assert hashlib.md5(original.read_bytes()).hexdigest() == '23c929502f692785881ce3dfcc167066'
sheet = load_workbook(original, read_only=True, data_only=False)['Table S3']
rows = {n: list(next(sheet.iter_rows(min_row=n, max_row=n, max_col=32, values_only=True)))
        for n in [1, 2, 3, 141, 143, 144, 145, 146]}
elements = rows[3][6:25]
assert elements == ['Sc','Mn','Fe','Rb','Cs','Ba','La','Ce','Nd','Sm','Eu','Tb','Dy','Yb','Lu','Hf','Ta','Th','U']
assert [rows[n][1] for n in [143,144,145,146]] == ['758-328','758-328','OH-2A','OH-2A']

def compare(a, b):
    values = []
    for i, element in enumerate(elements, 6):
        x, y = rows[a][i], rows[b][i]
        if not isinstance(x, (int, float)) or not isinstance(y, (int, float)) or x < 0 or y < 0:
            values.append({'element':element, 'a_raw':x, 'b_raw':y,
                           'symmetric_percent_difference':None})
        else:
            values.append({'element':element, 'a_raw':x, 'b_raw':y,
                           'symmetric_percent_difference':200*abs(x-y)/(x+y)})
    return {'rows':[a,b], 'values':values,
            'comparable_elements':sum(x['symmetric_percent_difference'] is not None for x in values)}

data = {'source_id':'S327', 'original_md5':hashlib.md5(original.read_bytes()).hexdigest(),
        'original_sha256':hashlib.sha256(original.read_bytes()).hexdigest(),
        'sheet':'Table S3', 'raw_rows':{str(n):v for n,v in rows.items()},
        'units':'Supplement: ppm except Fe percent; raw -1 means no data, not zero',
        'comparisons':{'cross_sample_D':compare(143,145), 'cross_sample_L':compare(144,146),
                       'same_sample_758_328_labs':compare(143,144),
                       'same_sample_OH_2A_labs':compare(145,146)},
        'rule':'200*absolute(a-b)/(a+b), descriptive only',
        'measurement_errors':None, 'correlation_threshold':None,
        'distinct_named_sample_count':2, 'physical_identity_authenticated':False,
        'correlation_accepted_by_this_audit':False,
        'FLV119_trace_element_join':None, 'age_transfer_verified':False}
(ROOT/'data/danville-inaa-comparison.json').write_text(json.dumps(data,indent=2)+'\n',encoding='utf-8')
print('Four analyses, two named samples; missing Dy excluded from affected comparisons.')
