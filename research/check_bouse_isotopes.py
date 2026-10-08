"""Reproduce S185 sample/genus oxygen-offset summaries; no age model or significance test.

Usage: python research/check_bouse_isotopes.py path/to/publisher-workbook.xlsx
Requires openpyxl and matplotlib. Input must match the audited publisher file.
"""
import hashlib
import json
import statistics
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

root = Path(__file__).resolve().parents[1]
path = Path(sys.argv[1])
expected = json.loads((root/'data/bouse-methods.json').read_text(encoding='utf-8'))['sha256']
assert hashlib.sha256(path.read_bytes()).hexdigest() == expected, 'Input differs from audited S185 version'
w = openpyxl.load_workbook(path.open('rb'), data_only=True)
sediments = {}
for row in range(6, 30):
    s = w['Supplement 1']
    sample = s.cell(row, 1).value
    assert sample not in sediments
    sediments[sample] = {'sample': sample, 'elevation_m_asl': s.cell(row,2).value,
        'unit': s.cell(row,3).value, 'oxygen_per_mil_vpdb': s.cell(row,19).value,
        'carbon_per_mil_vpdb': s.cell(row,20).value, 'locator': f'Supplement 1!A{row}:T{row}'}
groups = defaultdict(list)
records = []
for row in range(5, 58):
    s = w['Supplement 2']
    sample, elevation, unit, genus, oxygen, carbon = [s.cell(row,col).value for col in range(1,7)]
    sediment = sediments[sample]
    assert elevation == sediment['elevation_m_asl'] and unit == sediment['unit']
    assert isinstance(oxygen, (int,float)) and isinstance(carbon, (int,float))
    rec = {'sample':sample,'elevation_m_asl':elevation,'unit':unit,'genus':genus,
        'oxygen_per_mil_vpdb':oxygen,'carbon_per_mil_vpdb':carbon,
        'locator':f'Supplement 2!A{row}:F{row}', 'individual_specimen_id':None,
        'oxygen_minus_same_sample_micrite':round(oxygen-sediment['oxygen_per_mil_vpdb'],10)}
    records.append(rec)
    groups[(sample,genus)].append(rec)
summaries=[]
for (sample,genus), rows in groups.items():
    offsets=[r['oxygen_minus_same_sample_micrite'] for r in rows]
    summaries.append({'sample':sample,'genus':genus,'n_reported_rows':len(rows),
        'offset_mean':round(statistics.mean(offsets),6), 'offset_min':min(offsets),'offset_max':max(offsets),
        'offset_range':round(max(offsets)-min(offsets),6),
        'sediment_locator':sediments[sample]['locator'], 'ostracode_locators':[r['locator'] for r in rows]})
result={'source_id':'S185','input_sha256':expected,'review_status':'NOT_INDEPENDENTLY_REVIEWED',
    'method':'Subtract the same-sample micrite oxygen value from each reported ostracode oxygen value, then summarize within sample and genus. No pooling across samples or genera.',
    'limitations':['Reported rows lack individual specimen identifiers and do not establish independent biological replicates.',
        'All offsets in a sample share its single micrite value.', 'Means and ranges are descriptive, not confidence intervals.',
        'No age, salinity, temperature or hydrological mechanism estimated.'],
    'sediment_records':list(sediments.values()),'ostracode_records':records,'sample_genus_summaries':summaries}
(root/'analysis/bouse-isotope-offsets.json').write_text(json.dumps(result,indent=2,ensure_ascii=False)+'\n',encoding='utf-8')
fig, ax=plt.subplots(figsize=(10,8))
colors={'Cyprideis sp.':'#a14427','Cytheromorpha sp. A':'#246b9e','Candona spp.':'#407c47','Heterocypris sp.':'#864a99'}
for i, summary in enumerate(summaries):
    rows=groups[(summary['sample'],summary['genus'])]
    ax.scatter([r['oxygen_minus_same_sample_micrite'] for r in rows],[i]*len(rows),s=27,
        color=colors[summary['genus']],alpha=.75)
    ax.plot([summary['offset_min'],summary['offset_max']],[i,i],color=colors[summary['genus']],alpha=.4)
ax.set_yticks(range(len(summaries)),[s['sample']+' | '+s['genus']+' | n='+str(s['n_reported_rows']) for s in summaries])
ax.invert_yaxis()
ax.axvline(0,color='#777',linewidth=.8)
first_post=next(i for i,s in enumerate(summaries) if s['sample']=='HM 53')
ax.axhline(first_post-.5,color='#777',linestyle='--')
ax.set_xlabel('Ostracode minus same-sample micrite oxygen delta (per mil VPDB)')
ax.set_title('Hart Mine Wash: reported oxygen offsets by sample and genus\nAbove dashed line: pre-DCL marl; below: post-DCL marl and green claystone')
ax.grid(axis='x',alpha=.2)
fig.text(.02,.015,'S185, Supplement 1 / Supplement 2. Points are reported rows; overlaps can hide points. Lines show ranges, not confidence intervals.',fontsize=8)
fig.tight_layout(rect=(0,.04,1,1))
fig.savefig(root/'analysis/bouse-isotope-offsets.png',dpi=170)
plt.close(fig)
for s in summaries:
    print(f"{s['sample']} | {s['genus']} | {s['n_reported_rows']} | {s['offset_mean']:.2f} | {s['offset_min']:.1f} to {s['offset_max']:.1f}")
print(f'{len(sediments)} sediment samples; {len(records)} isotope rows; {len(summaries)} sample/genus groups; all joins match identifiers, elevations and units.')
